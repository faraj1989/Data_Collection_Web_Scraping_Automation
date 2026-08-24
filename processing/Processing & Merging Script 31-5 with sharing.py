import glob
import os
import re
import shutil
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from project_config import env_int, env_path

# --- CONFIGURATION ---
BASE_DIR = env_path("NOC_BASE_DIR", r"c:\Current_Alarms")
SHARED_FOLDER = env_path("NOC_SHARED_FOLDER", BASE_DIR / "Shared Current Alarms")
INTERVAL_SECONDS = env_int("MERGE_INTERVAL_SECONDS", 300)  # 5 minutes


# ========== ENSURE DIRECTORIES EXIST ==========
def ensure_dir(path):
    """Create directory if it doesn't exist."""
    if not path:
        return path
    if not os.path.exists(path):
        os.makedirs(path, exist_ok=True)
        print(f"📁 Created directory: {path}")
    return path


BASE_DIR = ensure_dir(BASE_DIR)
SHARED_FOLDER = ensure_dir(SHARED_FOLDER)


# =====================================================

def get_latest_date_folder():
    date_folders = [
        d for d in BASE_DIR.iterdir()
        if d.is_dir() and re.fullmatch(r"\d{4}-\d{2}-\d{2}", d.name)
    ]
    if not date_folders:
        return None
    return max(date_folders, key=os.path.getmtime)


def parse_mixed_dates(date_series):
    parsed_dates = []
    formats = [
        "%d/%m/%Y %H:%M:%S",
        "%d-%m-%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d-%b-%Y %H:%M:%S",
        "%m/%d/%Y %H:%M:%S %p",
        "%d/%m/%Y %H:%M:%S %p",
    ]

    for date_str in date_series:
        if pd.isna(date_str) or date_str == "-" or date_str == "":
            parsed_dates.append(pd.NaT)
            continue

        date_str = str(date_str).strip()
        parsed = pd.NaT

        for fmt in formats:
            try:
                parsed = pd.to_datetime(date_str, format=fmt)
                break
            except (ValueError, TypeError):
                continue

        if pd.isna(parsed):
            try:
                parsed = pd.to_datetime(date_str, dayfirst=True, errors="coerce")
            except Exception:
                parsed = pd.NaT

        parsed_dates.append(parsed)

    return pd.Series(parsed_dates)


def find_timestamp_column(df, source_name):
    timestamp_candidates = ["Last Occurred", "Last Occurred (NT)", "Last Occurred (ST)", "Occurrence Times"]

    for col in timestamp_candidates:
        if col in df.columns:
            non_null = df[col].ne("-").sum() if df[col].dtype == "object" else df[col].notna().sum()
            if non_null > 0:
                print(f"  {source_name}: Using '{col}' with {non_null} valid timestamps")
                return col

    for col in df.columns:
        if "occurred" in col.lower() or "time" in col.lower():
            non_null = df[col].ne("-").sum() if df[col].dtype == "object" else df[col].notna().sum()
            if non_null > 0:
                print(f"  {source_name}: Using discovered '{col}' with {non_null} valid timestamps")
                return col

    print(f"  WARNING: {source_name}: No valid timestamp column found")
    return None


def prepare_alarm_sheet(df, source_type):
    if df.empty:
        return pd.DataFrame(columns=["Severity", "Name", "Site Name", "Last Occurred", "First Occurred"])

    df_out = df.copy()

    if "Site Name" not in df_out.columns:
        if "MO Name" in df_out.columns:
            df_out["Site Name"] = df_out["MO Name"].astype(str).str.replace("(FN)", "", regex=False).str.strip()
        else:
            df_out["Site Name"] = "-"

    for col in ["Severity", "Name"]:
        if col not in df_out.columns:
            df_out[col] = "-"

    time_col = find_timestamp_column(df_out, source_type)
    if time_col:
        df_out["Last Occurred"] = df_out[time_col]
        df_out["_sort_time"] = parse_mixed_dates(df_out[time_col])
    else:
        df_out["Last Occurred"] = "-"
        df_out["_sort_time"] = pd.NaT

    first_occurred_col = None
    for col in ["First Occurred", "First Occurred (NT)", "First Occurred (ST)"]:
        if col in df_out.columns:
            first_occurred_col = col
            break

    df_out["First Occurred"] = df_out[first_occurred_col] if first_occurred_col else "-"
    df_out = df_out.sort_values("_sort_time", ascending=False, na_position="last")
    return df_out[["Severity", "Name", "Site Name", "Last Occurred", "First Occurred"]]


def read_latest_mae(latest_date_folder):
    mae_files = glob.glob(str(latest_date_folder / "CurrentAlarms_MAE_*.csv"))
    if not mae_files:
        raise FileNotFoundError(f"No MAE files found in {latest_date_folder}")

    latest_mae = max(mae_files, key=os.path.getmtime)
    df_mae_raw = pd.read_csv(latest_mae, skiprows=7).iloc[:, 1:]
    print(f"MAE file: {Path(latest_mae).name}")
    print("MAE columns:", df_mae_raw.columns.tolist())

    df_mae_raw = df_mae_raw[df_mae_raw["Name"].str.contains("NE Is Disconnected", na=False)].copy()

    mae_site_col_candidates = [
        col for col in df_mae_raw.columns
        if "site name" in col.lower() or "mo name" in col.lower()
    ]
    if not mae_site_col_candidates:
        raise ValueError("No site identifier column found in MAE file.")

    mae_site_col = mae_site_col_candidates[0]
    print(f"MAE: Using '{mae_site_col}' as site identifier")

    mae_time_col = find_timestamp_column(df_mae_raw, "MAE")
    if not mae_time_col:
        mae_time_col = "Last Occurred (NT)" if "Last Occurred (NT)" in df_mae_raw.columns else "Last Occurred"

    desired_cols_mae = ["Severity", "Name", mae_site_col, mae_time_col]
    existing_mae = [c for c in desired_cols_mae if c in df_mae_raw.columns]
    df_mae_raw = df_mae_raw[existing_mae].copy()
    df_mae_raw.rename(columns={mae_site_col: "MO Name", mae_time_col: "Last Occurred (NT)"}, inplace=True)
    return df_mae_raw


def read_latest_neteco(latest_date_folder):
    neteco_files = glob.glob(str(latest_date_folder / "CurrentAlarms_NetEco_*.csv"))
    if not neteco_files:
        print(f"WARNING: No NetEco files found in {latest_date_folder}")
        return 0, pd.DataFrame(columns=["Site Name", "Power Reason (NOC)", "Mains Failure Time"]), pd.DataFrame()

    latest_neteco = max(neteco_files, key=os.path.getmtime)
    df_ne_orig = pd.read_csv(latest_neteco, skiprows=7).iloc[:, 1:]
    print(f"NetEco file: {Path(latest_neteco).name}")
    print("NetEco columns:", df_ne_orig.columns.tolist())

    ne_site_col_candidates = [
        col for col in df_ne_orig.columns
        if "site name" in col.lower() or "mo name" in col.lower()
    ]
    if not ne_site_col_candidates:
        raise ValueError("No site identifier column found in NetEco file.")

    ne_site_col = ne_site_col_candidates[0]
    print(f"NetEco: Using '{ne_site_col}' as site identifier")

    ne_time_col = find_timestamp_column(df_ne_orig, "NetEco")
    if not ne_time_col:
        ne_time_col = "Last Occurred (NT)" if "Last Occurred (NT)" in df_ne_orig.columns else "Last Occurred"

    df_ne_orig.rename(columns={ne_site_col: "Site Name", ne_time_col: "Last Occurred"}, inplace=True)

    count_mains = df_ne_orig[df_ne_orig["Name"] == "Mains Failure"]["Site Name"].nunique()
    target_alarms = ["Mains Failure", "BLVD", "LLVD"]
    df_ne_filtered = df_ne_orig[df_ne_orig["Name"].isin(target_alarms)].copy()

    def format_power_row(group):
        found = list(group["Name"].unique())
        ordered_list = [a for a in target_alarms if a in found]
        reason_text = " - ".join(ordered_list)
        m_time = group[group["Name"] == "Mains Failure"]["Last Occurred"].head(1)
        m_time_str = m_time.iloc[0] if not m_time.empty else ""
        return pd.Series({"Power Reason (NOC)": reason_text, "Mains Failure Time": m_time_str})

    if df_ne_filtered.empty:
        df_ne_summary = pd.DataFrame(columns=["Site Name", "Power Reason (NOC)", "Mains Failure Time"])
    else:
        df_ne_summary = df_ne_filtered.groupby("Site Name", group_keys=False).apply(format_power_row).reset_index()

    return count_mains, df_ne_summary, df_ne_filtered


def apply_excel_formatting(output_path, df_final, df_metrics, df_neteco_sheet, df_mae_sheet, startcol_metrics):
    wb = load_workbook(output_path)

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    bold_font = Font(bold=True)

    def apply_formatting_to_sheet(ws, table_range, header_range=None):
        for row in ws[table_range]:
            for cell in row:
                cell.border = thin_border
        if header_range:
            for row in ws[header_range]:
                for cell in row:
                    cell.font = bold_font

    def auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    ws_sum = wb["Current Alarms Summary"]
    main_rows = df_final.shape[0] + 1
    main_cols = df_final.shape[1]
    if main_rows > 0 and main_cols > 0:
        main_range = f"A1:{get_column_letter(main_cols)}{main_rows}"
        apply_formatting_to_sheet(ws_sum, main_range, f"A1:{get_column_letter(main_cols)}1")

    metrics_rows = df_metrics.shape[0] + 1
    metrics_cols = df_metrics.shape[1]
    if metrics_rows > 0 and metrics_cols > 0:
        start_metrics_col = startcol_metrics + 1
        end_metrics_col = start_metrics_col + metrics_cols - 1
        metrics_range = f"{get_column_letter(start_metrics_col)}1:{get_column_letter(end_metrics_col)}{metrics_rows}"
        header_range = f"{get_column_letter(start_metrics_col)}1:{get_column_letter(end_metrics_col)}1"
        apply_formatting_to_sheet(ws_sum, metrics_range, header_range)

    auto_fit_columns(ws_sum)

    if not df_neteco_sheet.empty and "NetEco Current Alarms" in wb.sheetnames:
        ws_ne = wb["NetEco Current Alarms"]
        ne_rows = df_neteco_sheet.shape[0] + 1
        ne_cols = df_neteco_sheet.shape[1]
        ne_range = f"A1:{get_column_letter(ne_cols)}{ne_rows}"
        apply_formatting_to_sheet(ws_ne, ne_range, f"A1:{get_column_letter(ne_cols)}1")
        auto_fit_columns(ws_ne)

    ws_mae = wb["MAE Current Alarms"]
    mae_rows = df_mae_sheet.shape[0] + 1
    mae_cols = df_mae_sheet.shape[1]
    if mae_rows > 0 and mae_cols > 0:
        mae_range = f"A1:{get_column_letter(mae_cols)}{mae_rows}"
        apply_formatting_to_sheet(ws_mae, mae_range, f"A1:{get_column_letter(mae_cols)}1")
        auto_fit_columns(ws_mae)

    wb.save(output_path)


def process_latest_folder():
    global SHARED_FOLDER  # <-- Add this line at the start

    latest_date_folder = get_latest_date_folder()
    if not latest_date_folder:
        print("ERROR: No dated folders found under c:\\Current_Alarms.")
        return False

    print("=" * 70)
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Processing folder: {latest_date_folder}")

    df_mae_raw = read_latest_mae(latest_date_folder)
    count_disconnected = df_mae_raw["MO Name"].nunique()

    df_disc = df_mae_raw.copy()
    df_disc["Site Name"] = df_disc["MO Name"].astype(str).str.replace("(FN)", "", regex=False).str.strip()
    df_summary_mae = df_disc[["Site Name", "MO Name", "Name", "Last Occurred (NT)"]].drop_duplicates(subset="Site Name")
    df_summary_mae.rename(columns={"Name": "Alarm Name"}, inplace=True)

    count_mains, df_ne_summary, df_ne_filtered = read_latest_neteco(latest_date_folder)
    df_final = pd.merge(df_summary_mae, df_ne_summary, on="Site Name", how="left")

    if "Last Occurred (NT)" in df_final.columns:
        df_final.rename(columns={"Last Occurred (NT)": "Last Occurred"}, inplace=True)
    elif "Last Occurred" not in df_final.columns:
        df_final["Last Occurred"] = "-"

    if "Power Reason (NOC)" not in df_final.columns:
        df_final["Power Reason (NOC)"] = "Check TX/Link (No Power Alarm)"
    else:
        df_final["Power Reason (NOC)"] = df_final["Power Reason (NOC)"].fillna("Check TX/Link (No Power Alarm)")

    if "Mains Failure Time" not in df_final.columns:
        df_final["Mains Failure Time"] = "-"
    else:
        df_final["Mains Failure Time"] = df_final["Mains Failure Time"].fillna("-")

    required_columns = ["Site Name", "MO Name", "Alarm Name", "Last Occurred", "Power Reason (NOC)",
                        "Mains Failure Time"]
    for col in required_columns:
        if col not in df_final.columns:
            df_final[col] = "-"
    df_final = df_final[required_columns]

    df_metrics = pd.DataFrame({
        "Metric Description": [
            "Count of  NE Is Disconnected",
            "Count of Mains Failure",
        ],
        "Value": [count_disconnected, count_mains],
    })

    df_mae_sheet = prepare_alarm_sheet(df_mae_raw, "MAE")
    df_neteco_sheet = prepare_alarm_sheet(df_ne_filtered, "NetEco")

    report_time = datetime.now().strftime("%H%M")
    output_path = latest_date_folder / f"Final_NOC_Report_{latest_date_folder.name}_{report_time}.xlsx"

    # Ensure shared folder exists (using global)
    SHARED_FOLDER = ensure_dir(SHARED_FOLDER)
    shared_output_path = SHARED_FOLDER / "Live_NOC_Report.xlsx"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name="Current Alarms Summary", index=False, startrow=0, startcol=0)
        startcol_metrics = df_final.shape[1] + 2
        df_metrics.to_excel(
            writer,
            sheet_name="Current Alarms Summary",
            index=False,
            startrow=0,
            startcol=startcol_metrics,
        )

        if not df_neteco_sheet.empty:
            df_neteco_sheet.to_excel(writer, sheet_name="NetEco Current Alarms", index=False)
        df_mae_sheet.to_excel(writer, sheet_name="MAE Current Alarms", index=False)

    apply_excel_formatting(output_path, df_final, df_metrics, df_neteco_sheet, df_mae_sheet, startcol_metrics)
    print(f"Formatted report saved: {output_path}")

    try:
        shutil.copy2(output_path, shared_output_path)
        print(f"Shared Live Report updated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    except Exception as exc:
        print(f"WARNING: Could not update shared file, it might be open: {exc}")

    print(f"Done. Disconnected={count_disconnected}, Mains Failure={count_mains}")
    return True

def main():
    print("NOC report processor is running.")
    print(f"Input base folder: {BASE_DIR}")
    print(f"Loop interval: {INTERVAL_SECONDS // 60} minutes")

    while True:
        try:
            process_latest_folder()
        except Exception as exc:
            print(f"ERROR: Processing cycle failed: {exc}")

        print(f"Sleeping for {INTERVAL_SECONDS // 60} minutes. Press Ctrl+C to stop.")
        try:
            time.sleep(INTERVAL_SECONDS)
        except KeyboardInterrupt:
            print("Processor stopped by user.")
            break


if __name__ == "__main__":
    main()